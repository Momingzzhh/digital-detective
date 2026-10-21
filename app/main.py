import asyncio
import hashlib
import io
import json
import secrets
import sqlite3
import time
import uuid
import zipfile
from contextlib import asynccontextmanager
from pathlib import Path
from urllib.parse import urlparse

from fastapi import Depends, FastAPI, File, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, JSONResponse, PlainTextResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

from .agent import analyze
from .auth import hash_password, verify_password, session
from .config import ROOT, settings
from .local_model import available as local_available, adapter_path
from .store import initialize, db, save_case, load_case, budget_status

tasks = {}
jobs = {}
login_attempts = {}


@asynccontextmanager
async def lifespan(app):
    initialize()
    yield
    pending = list(tasks.values())
    for task in pending:
        task.cancel()
    await asyncio.gather(*pending, return_exceptions=True)


app = FastAPI(title='数字神探 · 研究工作台', lifespan=lifespan, docs_url=None, redoc_url=None)
app.mount('/static', StaticFiles(directory=ROOT / 'static'), name='static')
from .research import router as research_router
app.include_router(research_router)


@app.middleware('http')
async def protections(request, call_next):
    origin = request.headers.get('origin')
    if request.method not in ('GET', 'HEAD', 'OPTIONS') and origin and urlparse(origin).netloc != request.headers.get('host'):
        return JSONResponse({'detail': '拒绝跨站请求。'}, status_code=403)
    length = request.headers.get('content-length')
    if length and (not length.isdigit() or int(length) > 12 * 1024 * 1024):
        return JSONResponse({'detail': '请求过大。'}, status_code=413)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'"
    if request.url.path.startswith('/api'):
        response.headers['Cache-Control'] = 'no-store'
    return response


@app.get('/')
def home():
    return FileResponse(ROOT / 'static/index.html')


@app.get('/research')
def research_home():
    return FileResponse(ROOT / 'static/research.html')


@app.get('/health')
def health():
    return {'status': 'ok'}


class Credentials(BaseModel):
    name: str = Field(min_length=2, max_length=60)
    password: str = Field(min_length=10, max_length=200)


@app.post('/api/register')
def register(value: Credentials):
    if settings().get('ALLOW_REGISTRATION', 'true').lower() != 'true':
        raise HTTPException(403, '当前未开放注册。')
    with db() as con:
        try:
            con.execute('INSERT INTO users VALUES(?,?,?)', (uuid.uuid4().hex, value.name, hash_password(value.password)))
        except sqlite3.IntegrityError:
            raise HTTPException(409, '此用户名已存在。')
    return {'ok': True}


@app.post('/api/login')
def login(value: Credentials, request: Request, response: Response):
    host = request.client.host if request.client else 'unknown'
    attempts = [t for t in login_attempts.get(host, []) if time.time()-t < 60]
    if len(attempts) >= 10:
        raise HTTPException(429, '尝试过于频繁，请一分钟后重试。')
    login_attempts[host] = attempts + [time.time()]
    with db() as con:
        row = con.execute('SELECT * FROM users WHERE name=?', (value.name,)).fetchone()
        if not row or not verify_password(value.password, row['password']):
            raise HTTPException(401, '用户名或密码错误。')
        token, csrf = secrets.token_urlsafe(32), secrets.token_urlsafe(32)
        con.execute('DELETE FROM sessions WHERE expires<?', (time.time(),))
        con.execute('INSERT INTO sessions VALUES(?,?,?,?)', (hashlib.sha256(token.encode()).hexdigest(), row['id'], csrf, time.time()+86400))
    response.set_cookie('detective_session', token, httponly=True, samesite='strict', secure=settings().get('COOKIE_SECURE') == 'true', max_age=86400)
    return {'csrf': csrf, 'name': value.name}


@app.get('/api/me')
def me(user=Depends(session)):
    with db() as con:
        name = con.execute('SELECT name FROM users WHERE id=?', (user['user_id'],)).fetchone()[0]
    return {'name': name, 'csrf': user['csrf']}


@app.post('/api/logout')
def logout(response: Response, user=Depends(session)):
    with db() as con:
        con.execute('DELETE FROM sessions WHERE token=?', (user['token'],))
    response.delete_cookie('detective_session')
    return {'ok': True}


@app.get('/api/status')
def status(user=Depends(session)):
    cfg = settings()
    report = ROOT / 'data/dataset_report.json'
    five = ROOT / 'data/five_dimensions_report.json'
    five_status = json.loads(five.read_text()) if five.exists() else None
    if five_status:
        from .research import records, approved
        five_status['human_reviewed'] = sum(bool(approved(row)) for row in records())
    return {'model': '本地Qwen3 4B' + (' · 已配置适配器' if adapter_path() else ' · 基座模型'),
            'configured': local_available(), 'backend': 'local', 'external_api_enabled': False,
            'budget': budget_status(), 'limit': min(float(cfg.get('BUDGET_CNY', '500')), 500),
            'dataset': json.loads(report.read_text()) if report.exists() else {'qualified': 0, 'status': '案源核验中'},
            'five_dimensions': five_status,
            'registration': cfg.get('ALLOW_REGISTRATION', 'true') == 'true'}


class CaseInput(BaseModel):
    title: str = Field(min_length=1, max_length=120)
    text: str = Field(min_length=5, max_length=40000)


class MessageInput(BaseModel):
    text: str = Field(min_length=1, max_length=20000)


def owned(user, case_id):
    case = load_case(user['user_id'], case_id)
    if not case:
        raise HTTPException(404, '案件不存在。')
    return case


def assert_idle(case_id):
    if case_id in tasks and not tasks[case_id].done():
        raise HTTPException(409, '此案件正在分析，请等待本轮结束。')


@app.get('/api/cases')
def list_cases(user=Depends(session)):
    with db() as con:
        return [dict(r) for r in con.execute('SELECT id,title,updated FROM cases WHERE user_id=? ORDER BY updated DESC', (user['user_id'],))]


@app.post('/api/cases')
def create_case(value: CaseInput, user=Depends(session)):
    cid = uuid.uuid4().hex
    case = {'id': cid, 'title': value.title, 'sources': [{'id': 'S-1', 'name': '初始案情', 'text': value.text}],
            'messages': [], 'evidence': [], 'analysis': None, 'trace': [], 'graph': {'nodes': [], 'edges': []}}
    save_case(user['user_id'], case)
    return case


@app.get('/api/cases/{cid}')
def get_case(cid: str, user=Depends(session)):
    case = owned(user, cid)
    case['job'] = jobs.get(cid, {'state': 'idle'})
    return case


@app.delete('/api/cases/{cid}')
def delete_case(cid: str, user=Depends(session)):
    owned(user, cid)
    assert_idle(cid)
    with db() as con:
        con.execute('DELETE FROM cases WHERE id=? AND user_id=?', (cid, user['user_id']))
    jobs.pop(cid, None)
    return {'ok': True}


@app.post('/api/cases/{cid}/messages')
def message(cid: str, value: MessageInput, user=Depends(session)):
    case = owned(user, cid)
    assert_idle(cid)
    if case.get('research'):
        raise HTTPException(409, '验收会话请在五维复核页面释放原始证据；自由输入请新建普通案件。')
    if sum(len(s['text']) for s in case['sources']) + len(value.text) > 100000:
        raise HTTPException(413, '单案累计材料上限为 10 万字，请拆分案件。')
    case['messages'].append({'role': 'user', 'text': value.text, 'time': time.time()})
    case['sources'].append({'id': f'S-{len(case["sources"])+1}', 'name': '用户补充', 'text': value.text})
    save_case(user['user_id'], case)
    return case


@app.post('/api/cases/{cid}/analyze')
async def start_analysis(cid: str, user=Depends(session)):
    case = owned(user, cid)
    assert_idle(cid)
    if not local_available():
        raise HTTPException(503, '请先下载并校验本地模型后重试。')
    if sum(not t.done() for t in tasks.values()) >= 1:
        raise HTTPException(429, '本地模型已有分析任务运行，请稍后重试。')
    if case.get('research'):
        from .research import model_binding, get_record, approved
        if case['research']['model_binding'] != model_binding():
            raise HTTPException(409, '验收会话绑定的模型已改变。')
        row = get_record(case['research']['case_id'])
        if case['research']['record_binding'] != row['record_binding'] or not approved(row):
            raise HTTPException(409, '验收材料或人工复核状态已改变。')
        if case['research']['analysis_attempts'] >= 12:
            raise HTTPException(409, '验收每案最多12轮分析，请结束评分。')
        case['research']['analysis_attempts'] += 1
        save_case(user['user_id'], case)
    jobs[cid] = {'state': 'running', 'stage': '准备材料'}

    async def run():
        try:
            result = await analyze(case, lambda s: jobs[cid].update(stage=s))
            result['messages'].append({'role': 'assistant', 'text': result['analysis']['summary'] + '\n' + '\n'.join(result['analysis']['questions']), 'time': time.time()})
            save_case(user['user_id'], result)
            jobs[cid] = {'state': 'completed', 'stage': '分析完成'}
        except Exception as exc:
            from .llm import ModelError
            message = str(exc) if isinstance(exc, (ModelError, ValueError)) else '分析服务异常；已保存此前材料，请重试。'
            jobs[cid] = {'state': 'error', 'error': message}
        finally:
            tasks.pop(cid, None)
    tasks[cid] = asyncio.create_task(run())
    return jobs[cid]


@app.post('/api/extract-file')
async def extract_file(file: UploadFile = File(...), user=Depends(session)):
    raw = await file.read(10 * 1024 * 1024 + 1)
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(413, '文件最大 10 MB。')
    suffix = Path(file.filename or '').suffix.lower()
    try:
        if suffix in ('.txt', '.md'):
            text = raw.decode('utf-8-sig')
        elif suffix == '.docx':
            from docx import Document
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist()) > 40 * 1024 * 1024:
                    raise ValueError('解压尺寸过大')
            doc = Document(io.BytesIO(raw))
            text = '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in r.cells) for t in doc.tables for r in t.rows])
        elif suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            if len(reader.pages) > 100:
                raise ValueError('页数过多')
            text = '\n'.join(p.extract_text() or '' for p in reader.pages)
        else:
            raise HTTPException(415, '支持 TXT、Markdown、DOCX 和有文字层的 PDF。')
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, '文件无法解析；请使用 UTF-8 文本、DOCX 或有文字层的 PDF。')
    if len(text.strip()) < 5:
        raise HTTPException(422, '未提取到足够文字；扫描件请先转换为文字。')
    if len(text) > 40000:
        raise HTTPException(413, '提取文字超过 4 万字，请分批上传。')
    return {'name': Path(file.filename or '材料').name, 'text': text}


@app.get('/api/cases/{cid}/export')
def export(cid: str, user=Depends(session)):
    case = owned(user, cid)
    return Response(json.dumps(case, ensure_ascii=False, indent=2), media_type='application/json', headers={'Content-Disposition': f'attachment; filename="case-{cid}.json"'})
